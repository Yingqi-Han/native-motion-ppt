#!/usr/bin/env python3
"""Check click reveals of an existing PPT from bound actual frames, not animation XML alone."""
import argparse
import json
import math
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation

from check_native_motion import probe, frame_at
from check_native_playback import compare_frame
from native_motion.acceptance import sha, save
from native_motion.existing import inspect_existing, validate_plan, verify_changes
from native_motion.players import player_matches


def compare_region(actual, reference):
    # Never stretch a thin text line to 320x180: that amplifies a native one-pixel
    # baseline difference into a large mismatch. Keep the region's aspect ratio
    # before using the same existing 8/18 static descriptor thresholds.
    width=max(actual.width,math.ceil(actual.height*16/9));height=math.ceil(width*9/16)
    a=Image.new('RGB',(width,height),'#F4F2EA');b=a.copy()
    a.paste(actual,(0,0));b.paste(reference,(0,0))
    return compare_frame(a,b)


def measure_reveal(before, after, reference):
    initial, final = compare_region(before, reference), compare_region(after, reference)
    matches = final['rmse'] <= 8 and final['max_tile_rmse'] <= 18
    appeared = initial['rmse']-final['rmse'] >= 3 and initial['max_tile_rmse'] > 18
    return {'status': 'passed' if matches and appeared else 'failed', 'before_reference': initial,
            'after_reference': final, 'appearance_rmse_reduction': initial['rmse']-final['rmse'],
            'reason': 'Actual group appearance and complete reference match' if matches and appeared else
                      'Post-click content differs from the reference' if not matches else 'No reliable hidden-to-visible change at this click'}


def normalize_reference(reference, size):
    if reference.size == size:
        return reference, 'none'
    if reference.width == size[0] and abs(reference.height-size[1]) <= 1:
        return reference.resize(size, Image.Resampling.LANCZOS), 'one_pixel_PDF_raster_rounding'
    raise ValueError('Native reference and actual capture dimensions differ beyond PDF raster rounding')


def check(args):
    source, pptx, video = [Path(p).resolve() for p in (args.source, args.pptx, args.video)]
    plan = json.loads(Path(args.plan).read_text(encoding='utf-8'))
    inventory = inspect_existing(source);planned = validate_plan(plan, inventory)
    preservation = verify_changes(source, pptx, {s['part'] for s in inventory['slides'] if not s['has_animation']})
    proof = json.loads(Path(args.evidence).read_text(encoding='utf-8'))
    if not player_matches(args.player+'-desktop', proof):
        raise ValueError('Actual recording backend differs from the requested player')
    if proof.get('recording_sha256') != sha(video) or proof.get('source_pptx_sha256') != sha(pptx) or proof.get('opened_copy_sha256') != sha(pptx):
        raise ValueError('Stale recording/PPTX capture proof')
    interpretation = json.loads(Path(args.input_plan).read_text(encoding='utf-8'))
    if interpretation.get('playback_evidence_sha256') != sha(args.evidence) or interpretation.get('plan_sha256') != sha(args.plan):
        raise ValueError('Recorded key interpretation is bound to another capture or plan')
    actions = proof['actions']
    if not actions or actions[0]['key'] != 'F5' or any(a['key'] != 'space' for a in actions[1:]):
        raise ValueError('Expected one F5 followed by documented space actions')
    times = [a['at_wall_seconds'] for a in actions]
    if any(type(t) not in (int,float) or not math.isfinite(t) or t < 0 for t in times) or any(b <= a for a,b in zip(times,times[1:])):
        raise ValueError('Invalid actual input times')
    pages = interpretation['pages']
    if [p['slide'] for p in pages] != list(range(1, len(inventory['slides'])+1)):
        raise ValueError('Every original page must be associated in order')
    used = [];prs = Presentation(source)
    for page, original in zip(pages, inventory['slides']):
        indices = [page['page_action_index'], *[g['action_index'] for g in page['groups']]]
        if any(type(i) is not int or not 0 <= i < len(actions) for i in indices):
            raise ValueError('Invalid recorded action index')
        used += indices
        top = {o['id'] for o in original['objects']}
        if not page.get('context_ids') or any(type(i) is not int or i not in top for i in page['context_ids']):
            raise ValueError('Each page needs valid persistent context objects')
        if not original['has_animation']:
            if [g['shape_ids'] for g in page['groups']] != planned.get(page['slide'], []):
                raise ValueError('Recorded reveal groups differ from the applied plan')
        else:
            targets = {int(v) for v in prs.slides[page['slide']-1]._element.xpath('.//p:spTgt/@spid')}
            if any(not g['shape_ids'] or any(type(i) is not int or i not in top or i not in targets for i in g['shape_ids']) for g in page['groups']):
                raise ValueError('Preserved animation probes must target original animated objects')
        moving = {int(v) for v in prs.slides[page['slide']-1]._element.xpath('.//p:spTgt/@spid')}
        moving.update(i for g in page['groups'] for i in g['shape_ids'])
        if moving & set(page['context_ids']):
            raise ValueError('Context probes must remain static through this page')
    if used != list(range(len(actions))):
        raise ValueError('Key interpretation must account for every actual action exactly once and in order')
    meta, pts = probe(video);width,height = meta['streams'][0]['width'],meta['streams'][0]['height']
    if times[-1] >= pts[-1]:
        raise ValueError('Capture ends before its documented actions')
    out = Path(args.out).resolve();out.mkdir(parents=True,exist_ok=False)
    frames = [];contexts = [];reveals = [];reference_bindings = {};normalization = {}

    def sample(at, label):
        index = min(range(len(pts)), key=lambda i: abs(pts[i]-at));path = out/(label+'.png')
        image = frame_at(video,index,path)
        frames.append({'path':path.name,'timestamp':pts[index],'frame_index':index,'sha256':sha(path)})
        return image,frames[-1]

    def crop(image, ids, slide):
        objects={o['id']:o for o in inventory['slides'][slide-1]['objects']};boxes=[objects[i]['box_emu'] for i in ids]
        x=min(b[0] for b in boxes);y=min(b[1] for b in boxes);r=max(b[0]+b[2] for b in boxes);b=max(b[1]+b[3] for b in boxes)
        w,h=inventory['canvas_emu'];box=[max(0,math.floor(x/w*width)-2),max(0,math.floor(y/h*height)-2),min(width,math.ceil(r/w*width)+2),min(height,math.ceil(b/h*height)+2)]
        if box[2]<=box[0] or box[3]<=box[1]:raise ValueError('Probe has no visible area')
        return image.crop(box)

    settings=plan.get('animation',{})
    for pi,page in enumerate(pages):
        slide=page['slide'];path=Path(args.references)/f'slide-{slide:02}.png';reference=Image.open(path).convert('RGB')
        reference,normalization[path.name]=normalize_reference(reference,(width,height));reference_bindings[path.name]=sha(path)
        next_page=times[pages[pi+1]['page_action_index']] if pi+1<len(pages) else pts[-1]
        last_input=times[page['groups'][-1]['action_index']] if page['groups'] else times[page['page_action_index']]
        actual,location=sample(min(next_page-.35,last_input+1.4),f'page-{slide:02}-context')
        for sid in page['context_ids']:
            compared=compare_region(crop(actual,[sid],slide),crop(reference,[sid],slide))
            contexts.append({'slide':slide,'shape_id':sid,'status':'passed' if compared['rmse']<=8 and compared['max_tile_rmse']<=18 else 'failed','measurement':compared,'frame':location})
        for gi,group in enumerate(page['groups']):
            index=group['action_index'];at=times[index];ids=group['shape_ids']
            if inventory['slides'][slide-1]['has_animation']:settle=args.settle_seconds
            else:settle=(settings.get('duration_ms',450)+(len(ids)-1)*settings.get('stagger_ms',100))/1000+.4
            next_action=times[index+1] if index+1<len(times) else pts[-1]
            before,prior=sample(max(times[page['page_action_index']]+.5,at-.4),f'page-{slide:02}-group-{gi+1:02}-before')
            after,following=sample(min(at+settle,next_action-.25),f'page-{slide:02}-group-{gi+1:02}-after')
            measured=measure_reveal(crop(before,ids,slide),crop(after,ids,slide),crop(reference,ids,slide))
            if at+settle>next_action-.2:
                measured.update(status='pending',reason='Insufficient reading/settling interval before the next actual key; use a longer hold and new capture')
            reveals.append({'slide':slide,'group':gi+1,'shape_ids':ids,'original_animation':inventory['slides'][slide-1]['has_animation'],
                            'actual_key_seconds':at,'before':prior,'after':following,**measured})
    states=[r['status'] for r in contexts+reveals]
    status='failed' if 'failed' in states else 'pending' if 'pending' in states else 'passed'
    report={'format_version':1,'origin':'measurement_of_actual_desktop_recording','status':status,
            'recording_sha256':sha(video),'source_pptx_sha256':sha(pptx),'original_pptx_sha256':sha(source),
            'plan_sha256':sha(args.plan),'playback_evidence_sha256':sha(args.evidence),'input_interpretation_sha256':sha(args.input_plan),
            'player':proof['player'],'version':proof['version'],'preservation':preservation,'contexts':contexts,'reveals':reveals,
            'reference_sha256':reference_bindings,'reference_normalization':normalization,'frames':frames,
            'thresholds':{'held_reference_rmse':8,'held_max_tile_rmse':18,'appearance_rmse_reduction':3},
            'comparison_geometry':'Aspect-preserving 16:9 padding before the existing static descriptor; thin crops are never stretched',
            'scope':'Original key indices define conservative before/after probes; original video PTS selects actual frames. Checks whole-object appearance and persistent context only, not exact onset, arbitrary motion paths, native chart data, Morph or automatic advance.'}
    sheet=Image.new('RGB',(1600,math.ceil(len(frames)/3)*324),'#102326');draw=ImageDraw.Draw(sheet)
    for i,frame in enumerate(frames):
        image=Image.open(out/frame['path']);image.thumbnail((532,299));x=i%3*534;y=i//3*324;sheet.paste(image,(x,y));draw.text((x+8,y+300),f"{frame['timestamp']:.3f}s / frame {frame['frame_index']}",fill='#FFFFFF')
    sheet.save(out/'contact.jpg',quality=93);save(out/'enhancement-playback-audit.json',report)
    return report


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('video')
    for name in ['source','pptx','plan','evidence','references','input-plan','out']:p.add_argument('--'+name,required=True)
    p.add_argument('--player',choices=['wps','impress'],default='wps');p.add_argument('--settle-seconds',type=float,default=1.2)
    args=p.parse_args()
    if not math.isfinite(args.settle_seconds) or not .3<=args.settle_seconds<=8:p.error('settle-seconds must be between .3 and 8')
    r=check(args);print(json.dumps({'status':r['status'],'context_objects':len(r['contexts']),'reveals':len(r['reveals']),'scope':r['scope']},indent=2))
    raise SystemExit(1 if r['status']=='failed' else 2 if r['status']=='pending' else 0)
