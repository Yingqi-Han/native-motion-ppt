#!/usr/bin/env python3
"""Verify an isolated two-slide rectangle Morph using actual capture pixels and PTS."""
import argparse
import json
import math
from pathlib import Path
from zipfile import ZipFile

import numpy as np
from lxml import etree
from pptx import Presentation

from check_native_motion import probe, decode
from native_motion.acceptance import sha, save
from native_motion.playback_motion import colored_box


def check(args):
    if not all(math.isfinite(v) for v in [args.start,args.end]) or not 0<=args.start<args.end:
        raise ValueError('Use a finite nonnegative transition window')
    video=Path(args.video).resolve();pptx=Path(args.pptx).resolve()
    evidence=json.loads(Path(args.evidence).read_text())
    if evidence.get('origin')!='native_desktop_screen_capture' or evidence.get('is_pptx_playback_recording') is not True:
        raise ValueError('A real desktop recording is required')
    if evidence['recording_sha256']!=sha(video) or evidence['source_pptx_sha256']!=sha(pptx):
        raise ValueError('Capture/PPTX evidence checksum mismatch')
    prs=Presentation(pptx)
    if len(prs.slides)!=2 or any(len(s.shapes)!=1 for s in prs.slides):
        raise ValueError('Use exactly two slides, each with one isolated solid rectangle')
    a,b=[s.shapes[0] for s in prs.slides]
    if a.name!=b.name or not a.name.startswith('!!') or a.fill.fore_color.rgb!=b.fill.fore_color.rgb:
        raise ValueError('Same explicit Morph identity and solid color required')
    ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main',
        'p159':'http://schemas.microsoft.com/office/powerpoint/2015/09/main'}
    with ZipFile(pptx) as z:
        roots=[etree.fromstring(z.read(f'ppt/slides/slide{i}.xml')) for i in [1,2]]
        if any(r.xpath('//p:timing',namespaces=ns) for r in roots):
            raise ValueError('Diagnostic must contain no object animations')
        if not roots[1].xpath('//p159:morph',namespaces=ns):raise ValueError('Native Morph transition is missing')
    meta,times=probe(video);stream=meta['streams'][0];vw,vh=stream['width'],stream['height']
    boxes=[[s.left*vw/prs.slide_width,s.top*vh/prs.slide_height,
            (s.left+s.width)*vw/prs.slide_width,(s.top+s.height)*vh/prs.slide_height] for s in [a,b]]
    x0=max(0,int(min(x[0] for x in boxes))-15);y0=max(0,int(min(x[1] for x in boxes))-15)
    x1=min(vw,int(max(x[2] for x in boxes))+16);y1=min(vh,int(max(x[3] for x in boxes))+16)
    indices=[i for i,t in enumerate(times) if args.start<=t<=args.end]
    if len(indices)<12:raise ValueError('Select a window containing both holds and the entire transition')
    filt=f'select=between(n\\,{indices[0]}\\,{indices[-1]}),crop={x1-x0}:{y1-y0}:{x0}:{y0}:exact=1'
    records=[];color=list(bytes(a.fill.fore_color.rgb))
    for i,t,pixels in decode(video,filt,(x1-x0,y1-y0),indices,times,end=args.end):
        found=colored_box(pixels,color,tolerance=24)
        box=None;occupancy=0
        if found:
            local,area=found;box=[local[0]+x0,local[1]+y0,local[2]+x0,local[3]+y0]
            occupancy=area/((local[2]-local[0])*(local[3]-local[1]))
        records.append({'timestamp':t,'frame_index':i,'box':box,'solid_occupancy':occupancy})
    visible=[r for r in records if r['box']]
    def near(actual,expected):return max(abs(x-y) for x,y in zip(actual,expected))<=3
    initial=[r for r in visible if near(r['box'],boxes[0])]
    final=[r for r in visible if near(r['box'],boxes[1])]
    middle=[r for r in visible if not near(r['box'],boxes[0]) and not near(r['box'],boxes[1])]
    errors=[]
    if len(initial)<6 or len(final)<6:errors.append('Both endpoint holds were not captured')
    centers=[(box[0]+box[2])/2 for box in boxes]
    valid=[r for r in middle if r['solid_occupancy']>=.95 and centers[0]+1<(r['box'][0]+r['box'][2])/2<centers[1]-1
           and min(a.width,b.width)*vw/prs.slide_width-3<=r['box'][2]-r['box'][0]<=max(a.width,b.width)*vw/prs.slide_width+3]
    poses=len(set(tuple(r['box']) for r in valid))
    if poses<10:errors.append('Fewer than ten solid intermediate poses: no proof of continuous native Morph')
    if len(valid)!=len(middle):errors.append('Non-solid or geometrically invalid intermediate pose: inspect for fade/double image')
    seq=[r for r in visible if initial and final and initial[-1]['timestamp']<=r['timestamp']<=final[0]['timestamp']]
    coords=np.array([r['box'] for r in seq],dtype=float)
    max_step=float(np.abs(np.diff(coords,axis=0)).max()) if len(coords)>1 else None
    max_gap=max((y['timestamp']-x['timestamp'] for x,y in zip(seq,seq[1:])),default=0)
    if max_step is None or max_step>90:errors.append('Large frame-to-frame jump in isolated Morph')
    if max_gap>.1:errors.append('Capture gap greater than 100ms during Morph')
    if len(coords)>1 and np.diff(coords[:,0]).min()<-2:errors.append('Unexpected backward jump')
    report={'status':'failed' if errors else 'passed','origin':'measurement_of_actual_desktop_recording',
            'player':evidence['player'],'recording_sha256':sha(video),'source_pptx_sha256':sha(pptx),
            'native_transition_present':True,'object_animations_present':False,'endpoint_boxes_px':boxes,
            'solid_intermediate_distinct_poses':poses,'max_edge_step_px':max_step,'max_capture_gap_seconds':max_gap,
            'observed_transition_seconds':[seq[0]['timestamp'],seq[-1]['timestamp']] if seq else None,
            'errors':errors,'records':records,
            'scope':'Isolated rectangle, left-to-right position and scale Morph in this exact recorded build. No source animation evaluator; fade and double-image frames cannot pass. Complex group Morph requires its own visual review.'}
    save(Path(args.out),report);return report


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('video');p.add_argument('--pptx',required=True)
    p.add_argument('--evidence',required=True);p.add_argument('--start',type=float,required=True);p.add_argument('--end',type=float,required=True);p.add_argument('--out',required=True)
    r=check(p.parse_args());print(json.dumps({k:v for k,v in r.items() if k!='records'},indent=2))
    raise SystemExit(1 if r['errors'] else 0)
