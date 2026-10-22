#!/usr/bin/env python3
"""Observe playback activity and track isolated native objects from real capture."""
import argparse
import json
import math
import subprocess
from pathlib import Path

import numpy as np
from PIL import Image,ImageFilter,ImageDraw
from pptx import Presentation

from native_motion.acceptance import save,sha
from native_motion.model import load
from native_motion.playback_motion import analyse_trace,colored_box,foreground_box,intervals,stable_runs,thumbnail_score


def probe(video):
    data=json.loads(subprocess.check_output(["ffprobe","-v","error","-select_streams","v:0","-show_entries",
         "stream=width,height,avg_frame_rate,nb_frames:format=duration:frame=best_effort_timestamp_time","-of","json",str(video)],text=True))
    times=[float(f["best_effort_timestamp_time"]) for f in data["frames"]]
    if len(times)<6 or any(not math.isfinite(t) for t in times) or any(b<=a for a,b in zip(times,times[1:])):
        raise ValueError("Recording requires strictly increasing actual PTS")
    return data,times


def decode(video,filter_text,size,indices,times,end=None):
    command=["ffmpeg","-hide_banner","-loglevel","error","-threads","2","-i",str(video)]
    if end is not None:command +=["-to",str(end+.25)]
    command +=["-vf",filter_text,"-fps_mode","passthrough","-f","rawvideo","-pix_fmt","rgb24","pipe:1"]
    process=subprocess.Popen(command,stdout=subprocess.PIPE);length=size[0]*size[1]*3;count=0
    try:
        for index in indices:
            chunks=[];remaining=length
            while remaining:
                raw=process.stdout.read(remaining)
                if not raw:raise RuntimeError("Decoded capture ended before its PTS list")
                chunks.append(raw);remaining-=len(raw)
            count+=1
            yield index,times[index],np.frombuffer(b"".join(chunks),dtype=np.uint8).reshape(size[1],size[0],3)
        if process.stdout.read(1):raise RuntimeError("Decoded capture has extra frames without PTS")
        if process.wait()!=0:raise RuntimeError("Video decode failed")
    finally:
        process.stdout.close()
        if process.poll() is None:process.terminate();process.wait()


def frame_at(video,index,path):
    subprocess.run(["ffmpeg","-hide_banner","-loglevel","error","-y","-threads","2","-i",str(video),
          "-vf",f"select=eq(n\\,{index})","-frames:v","1",str(path)],check=True)
    return Image.open(path).convert("RGB")


def measure_bar(frame,color):
    full=colored_box(frame,color)
    if not full:return None
    box,area=full;middle=(box[0]+box[2])//2;a=max(0,middle-5);b=min(frame.shape[1],middle+5)
    central=colored_box(frame[:,a:b],color)
    if not central:return None
    bounds,_=central
    return [box[0],bounds[1],box[2],bounds[3]],area


def tighten_actor_roi(obj,reference):
    """Use the component nearest the object's static center, excluding nearby labels."""
    a,b,c,d=obj["roi"];crop=reference[b:d,a:c]
    border=np.concatenate([crop[0],crop[-1],crop[:,0],crop[:,-1]])
    background=np.median(border,axis=0).astype(int)
    mask=np.max(np.abs(crop.astype(np.int16)-background),axis=2)>38
    ys,xs=np.nonzero(mask)
    if len(xs)<16:return
    x,y,w,h=obj["source_box"];cx=x+w/2-a;cy=y+h/2-b
    nearest=int(np.argmin((xs-cx)**2+(ys-cy)**2))
    component=Image.fromarray(mask.astype(np.uint8)*255).copy()
    ImageDraw.floodfill(component,(int(xs[nearest]),int(ys[nearest])),128,thresh=0)
    connected=np.array(component)==128;yy,xx=np.nonzero(connected)
    if len(xx)<16:return
    dx0,dy0,dx1,dy1=obj["search_offsets"]
    obj["roi"]=[max(a,math.floor(a+xx.min()-max(abs(dx0),abs(dx1))-6)),max(b,math.floor(b+yy.min()-max(abs(dy0),abs(dy1))-6)),
                min(c,math.ceil(a+xx.max()+1+max(abs(dx0),abs(dx1))+6)),min(d,math.ceil(b+yy.max()+1+max(abs(dy0),abs(dy1))+6))]


def eligible(doc,prs):
    objects=[];uncovered=[]
    for si,slide in enumerate(doc["slides"],1):
        shapes={s.name.lstrip("!"):s for s in prs.slides[si-1].shapes}
        for obj in slide["objects"]:
            effects=obj["animations"]
            if not effects:continue
            bar=obj["kind"]=="rect" and obj.get("role")=="data" and any(a["kind"]=="enter" and a.get("effect")=="wipe" and a["direction"]=="up" for a in effects)
            mover=any(a["kind"]=="move" for a in effects)
            transform=any(a["kind"] in {"scale","rotate"} for a in effects)
            if not(bar or mover or transform):continue
            reason=None
            if transform:reason="缩放或旋转需要模板尺度/角度追踪，本版只记录其待检查状态"
            elif obj["opacity"]!=1 or obj["rotation"]:reason="透明或旋转对象不能可靠地用当前轮廓标定"
            elif bar and (obj.get("gradient") or obj["fill"]=="none" or mover):reason="数据柱包含渐变或并行位移，需要单独标定"
            elif obj["kind"] not in {"rect","ellipse","group"}:reason="当前追踪器支持隔离矩形、椭圆与组合几何"
            if reason:
                code='unsupported-transform' if transform else 'unsupported-opacity-rotation' if obj['opacity']!=1 or obj['rotation'] else 'unsupported-bar-composition' if bar else 'unsupported-geometry'
                uncovered.append({"slide":si,"object":obj["id"],"status":"pending","reason":reason,'reason_code':code});continue
            shape=shapes.get(obj["id"])
            if shape is None:raise ValueError("Source object missing from PPTX: "+obj["id"])
            actual=[shape.left*doc["canvas"]["width"]/prs.slide_width,shape.top*doc["canvas"]["height"]/prs.slide_height,
                    shape.width*doc["canvas"]["width"]/prs.slide_width,shape.height*doc["canvas"]["height"]/prs.slide_height]
            if max(abs(a-b) for a,b in zip(actual,obj["box"]))>.1:raise ValueError("Source/PPT geometry mismatch: "+obj["id"])
            offsets=[[0,0]]
            for effect in effects:
                offsets+=effect.get("path",[effect.get("from",[0,0]),effect.get("to",[0,0])])
            x,y,w,h=obj["box"];padding=24
            roi=[max(0,math.floor(x+min(p[0] for p in offsets)-padding)),max(0,math.floor(y+min(p[1] for p in offsets)-padding)),
                 min(doc["canvas"]["width"],math.ceil(x+w+max(p[0] for p in offsets)+padding)),min(doc["canvas"]["height"],math.ceil(y+h+max(p[1] for p in offsets)+padding))]
            paths=[a.get("path",[a.get("from",[0,0]),a.get("to",[0,0])]) for a in effects if a["kind"]=="move"]
            intent="closed" if paths and all(p[0]==p[-1] for p in paths) else "linear" if len(paths)==1 and len(paths[0])==2 else "free"
            objects.append({"slide":si,"object":obj["id"],"pptx_shape_id":shape.shape_id,"kind":"bar" if bar else "actor", "intent":intent,
                            "roi":roi,"source_box":obj["box"],"search_offsets":[min(p[0] for p in offsets),min(p[1] for p in offsets),max(p[0] for p in offsets),max(p[1] for p in offsets)],
                            "color":list(bytes.fromhex(obj["fill"])) if bar else None})
    return objects,uncovered


def observe_pages(video,times,plan,references,out):
    refs=[np.array(Image.open(Path(references)/f"slide-{p['slide']:02}.png").convert("RGB").resize((320,180)).filter(ImageFilter.GaussianBlur(1.2))) for p in plan["pages"]]
    activity=[];scores=[[] for _ in refs];previous=None
    for index,t,array in decode(video,"scale=320:180:flags=area",(320,180),range(len(times)),times):
        blur=np.array(Image.fromarray(array).filter(ImageFilter.GaussianBlur(1.2)))
        value=thumbnail_score(blur,previous)[1] if previous is not None else 0
        activity.append({"timestamp":t,"frame_index":index,"max_tile_change_rmse":value})
        previous=blur
        for pi,p in enumerate(plan["pages"]):
            if p["page_input_seconds"]-.15<=t<p["next_page_seconds"]-.20:
                global_error,tile=thumbnail_score(blur,refs[pi]);scores[pi].append((index,global_error,tile))
    events=intervals(times,[a["max_tile_change_rmse"]>1.2 for a in activity])
    observed=json.loads(json.dumps(plan));observed["origin"]="measurement_of_actual_desktop_recording"
    pages=[];landmarks=[]
    for pi,p in enumerate(observed["pages"]):
        rows=scores[pi];local_times=[times[r[0]] for r in rows]
        valid=[r[1]<=8 and r[2]<=18 and activity[r[0]]["max_tile_change_rmse"]<=1.2
               and times[r[0]]>max(p["click_at_seconds"],default=p["page_input_seconds"])+.2 for r in rows]
        runs=stable_runs(local_times,valid)
        if not runs:
            pages.append({"slide":p["slide"],"status":"pending","reason":"没有找到足够长且与该页参考一致的稳定画面",'reason_code':'page-not-calibrated'});continue
        a,b=runs[-1];pick=max(a,b-round(.35/max(.001,np.median(np.diff(local_times)))))
        index=rows[pick][0];p["held_probe_seconds"]=times[index]
        # Broad input windows assign page identity; actual activity supplies landmarks.
        local=[e for e in events if e["end"]>=p["page_input_seconds"]-.15 and e["start"]<p["next_page_seconds"]-.2]
        for e in local:landmarks +=[e["start"],e["end"]]
        pages.append({"slide":p["slide"],"status":"passed","stable_interval_seconds":[local_times[a],local_times[b]],
                      "held_timestamp":times[index],"held_frame_index":index,"activity_intervals":local,
                      "reference_rmse":rows[pick][1],"reference_tile_rmse":rows[pick][2]})
    observed["observed_landmarks"]=sorted(set(landmarks))
    return activity,events,pages,observed


def check(args):
    video=Path(args.video).resolve();pptx=Path(args.pptx).resolve();out=Path(args.out).resolve()
    evidence=json.loads(Path(args.evidence).read_text())
    if evidence.get("origin")!="native_desktop_screen_capture" or evidence.get("is_pptx_playback_recording") is not True:
        raise ValueError("Only actual desktop screen capture is accepted; design previews are not measurements")
    if evidence["recording_sha256"]!=sha(video) or evidence["source_pptx_sha256"]!=sha(pptx):raise ValueError("Recording or PPTX differs from evidence")
    doc=load(args.spec);prs=Presentation(pptx);plan=json.loads(Path(args.plan).read_text());meta,times=probe(video)
    phase_path=getattr(args,'phases',None);phase=None
    if phase_path:
        phase=json.loads(Path(phase_path).read_text())
        if phase.get('origin')!='measurement_of_actual_desktop_recording' or any(phase.get(k)!=v for k,v in {
            'recording_sha256':sha(video),'source_pptx_sha256':sha(pptx),'source_json_sha256':sha(args.spec)}.items()):raise ValueError('Phase evidence belongs to another source/recording')
        if any(plan.get(k)!=phase.get(k) for k in ['recording_sha256','source_pptx_sha256','source_json_sha256']):raise ValueError('Observed plan does not bind the phase evidence')
        if phase.get('format_version',1)>=2 and any(d.get('playback_evidence_sha256')!=sha(args.evidence) for d in [phase,plan]):raise ValueError('Phase/plan input evidence belongs to another capture')
    stream=meta["streams"][0];vw,vh=stream["width"],stream["height"]
    if (vw,vh)!=(doc["canvas"]["width"],doc["canvas"]["height"]):raise ValueError("Motion measurement requires capture and source canvas dimensions to match; no letterbox inference")
    if len(prs.slides)!=len(doc["slides"]) or [p["slide"] for p in plan["pages"]]!=list(range(1,len(prs.slides)+1)):
        raise ValueError("Source, playback plan and PPTX page identities differ")
    if any(p["next_page_seconds"]>=times[-1] for p in plan["pages"]):raise ValueError("Recording does not cover the complete planned playback")
    if not math.isfinite(args.max_step) or args.max_step<=0:raise ValueError("max-step must be positive")
    out.mkdir(parents=True,exist_ok=True)
    candidates,uncovered=eligible(doc,prs)
    activity,events,pages,observed=observe_pages(video,times,plan,args.references,out)
    tracked=[];errors=[];landmarks=observed["observed_landmarks"]
    for page in plan["pages"]:
        objects=[o for o in candidates if o["slide"]==page["slide"]]
        if not objects:continue
        identified=pages[page["slide"]-1]
        if identified["status"]!="passed":
            uncovered +=[{"slide":page["slide"],"object":o["object"],"status":"pending","reason":"该页尚未取得可信的实际到位标定",'reason_code':'page-not-calibrated'} for o in objects];continue
        reference_path=out/f"held-{page['slide']:02}.png"
        reference=np.array(frame_at(video,identified["held_frame_index"],reference_path))
        for obj in objects:
            if obj["kind"]=="actor":tighten_actor_roi(obj,reference)
            x0,y0,x1,y1=obj["roi"];crop=reference[y0:y1,x0:x1]
            border=np.concatenate([crop[0],crop[-1],crop[:,0],crop[:,-1]])
            obj["background"]=np.median(border,axis=0).astype(int).tolist()
            x,y,w,h=obj["source_box"];obj["seed"]=[x+w/2-x0,y+h/2-y0]
            calibration=measure_bar(crop,obj["color"]) if obj["kind"]=="bar" else foreground_box(crop,obj["background"],trim_border=4)
            obj["calibration"]=calibration;obj["records"]=[]
        x0=min(o["roi"][0] for o in objects);y0=min(o["roi"][1] for o in objects)
        x1=max(o["roi"][2] for o in objects);y1=max(o["roi"][3] for o in objects)
        start=max(0,page["page_input_seconds"]-.15);end=page["next_page_seconds"]-.20
        indices=[i for i,t in enumerate(times) if start<=t<end]
        filter_text=f"select=between(n\\,{indices[0]}\\,{indices[-1]}),crop={x1-x0}:{y1-y0}:{x0}:{y0}:exact=1"
        for index,t,array in decode(video,filter_text,(x1-x0,y1-y0),indices,times,end=end):
            for obj in objects:
                a,b,c,d=obj["roi"];crop=array[b-y0:d-y0,a-x0:c-x0]
                measured=measure_bar(crop,obj["color"]) if obj["kind"]=="bar" else foreground_box(crop,obj["background"],trim_border=4)
                calibration=obj["calibration"];box=None
                if measured and calibration:
                    bounds,area=measured;ref,refarea=calibration
                    if obj["kind"]=="bar":ok=bounds[3]-bounds[1]>=3 and abs((bounds[2]-bounds[0])-(ref[2]-ref[0]))<=3
                    else:
                        ok=.85<=area/refarea<=1.15 and all(abs((bounds[k+2]-bounds[k])-(ref[k+2]-ref[k]))<=2 for k in [0,1])
                        ok=ok and bounds[0]>1 and bounds[1]>1 and bounds[2]<c-a-1 and bounds[3]<d-b-1
                    if ok:box=[bounds[0]+a,bounds[1]+b,bounds[2]+a,bounds[3]+b]
                obj["records"].append({"timestamp":t,"frame_index":index,"box":box})
        for obj in objects:
            raw=analyse_trace(obj["records"],obj["kind"],obj["intent"],args.max_step)
            cutoff=page.get('object_motion_starts',{}).get(obj['object'],page.get('object_motion_start_seconds')) if obj['kind']=='actor' else None
            measurements=obj['records'];phase_info=None
            if cutoff is not None:
                incoming=[m for m in (phase or {}).get('morphs',[]) if m['slide']==obj['slide'] and m['object']==obj['object']]
                if not incoming or any(m['status']!='passed' or m['settled_seconds']>cutoff+.001 for m in incoming):raise ValueError('A phase cutoff requires a passed measured Morph arrival for this object')
                if not page['page_input_seconds']<=cutoff<page['next_page_seconds']:raise ValueError('Invalid object phase cutoff')
                measurements=[r for r in measurements if r['timestamp']>=cutoff]
                phase_info={'start_seconds':cutoff,'evidence_sha256':sha(phase_path),'method':'measured_incoming_morph_arrival_plateau'}
            measured=analyse_trace(measurements,obj["kind"],obj["intent"],args.max_step)
            records=obj.pop("records");obj.pop("calibration")
            row={**obj,**measured,"records":records,"decoded_frames":len(records),"held_frame_index":identified["held_frame_index"]}
            row.update(raw_status=raw['status'],raw_issues=raw['issues'],phase=phase_info)
            tracked.append(row)
            for issue in row["issues"]:
                issue.update(slide=obj["slide"],object=obj["object"],pptx_shape_id=obj["pptx_shape_id"])
                if issue["severity"]=="failed":errors.append(issue)
            for event in row["observed_intervals"]:landmarks +=[event["start"],event["end"]]
            if row.get("kinematics"):landmarks.append(times[row["peak_frame_index"]])
    observed.update(recording_sha256=sha(video),source_pptx_sha256=sha(pptx),source_json_sha256=sha(args.spec),
                    planned_input_sha256=sha(args.plan),observed_landmarks=sorted(set(landmarks)))
    save(out/"observed-playback-plan.json",observed)
    status="failed" if errors else "pending" if uncovered or any(p["status"]!="passed" for p in pages) or any(o["status"]!="passed" for o in tracked) else "passed"
    report={"format_version":1,"origin":"measurement_of_actual_desktop_recording","recording_sha256":sha(video),"source_pptx_sha256":sha(pptx),
            "source_json_sha256":sha(args.spec),"planned_input_sha256":sha(args.plan),"status":status,"frames_decoded_for_activity":len(times),
            "capture_metadata":{k:v for k,v in meta.items() if k!="frames"},"pages":pages,"activity_intervals":events,"activity":activity,
            "objects":tracked,"uncovered_objects":uncovered,"errors":errors,"thresholds":{"max_actor_step_px":args.max_step,"baseline_px":1,"linear_rebound_px":2,"interior_hold_seconds":.30},
            "scope":"Actual decoded pixels and PTS, no inserted frames. Planned keys associate page windows; source paths only bound search and describe intent. Foreground tracking requires isolated geometry, not arbitrary overlapping shapes. Activity/holds are pixel observations, not proof of player internal timing or native Morph."}
    report['phase_evidence_sha256']=sha(phase_path) if phase_path else None
    save(out/"native-motion-playback-audit.json",report)
    from native_motion.motion_report import write_motion_report
    write_motion_report(out,report,video)
    return report


if __name__=="__main__":
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("video");parser.add_argument("--pptx",required=True);parser.add_argument("--evidence",required=True)
    parser.add_argument("--spec",required=True);parser.add_argument("--plan",required=True);parser.add_argument("--references",required=True);parser.add_argument("--out",required=True)
    parser.add_argument("--max-step",type=float,default=24)
    parser.add_argument('--phases',help='SHA-bound measured phase audit; permits local drift to be analysed after the observed incoming Morph')
    result=check(parser.parse_args())
    print(json.dumps({"status":result["status"],"activity_intervals":len(result["activity_intervals"]),"tracked_objects":len(result["objects"]),
                      "uncovered_objects":result["uncovered_objects"],"errors":result["errors"]},ensure_ascii=False,indent=2))
    if result["errors"]:raise SystemExit(1)
