#!/usr/bin/env python3
"""Measure bottom edges and upward reveals from actual player-recording pixels."""
import argparse
import hashlib
import json
import math
import subprocess
from pathlib import Path

from PIL import Image,ImageChops
from pptx import Presentation


def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def color_bounds(image,color,tolerance=34):
    difference=ImageChops.difference(image,Image.new("RGB",image.size,color))
    r,g,b=difference.split()
    mask=ImageChops.lighter(ImageChops.lighter(r,g),b).point(lambda value:255 if value<=tolerance else 0)
    return mask.getbbox()


def check(args):
    if not all(math.isfinite(t) for t in [args.start,args.end,args.hold,args.tolerance]) or not 0<=args.start<args.end or args.hold<0 or args.tolerance<0:
        raise ValueError("Provide valid recording timestamps")
    video=Path(args.video).resolve();deck=Path(args.pptx).resolve();out=Path(args.out).resolve()
    proof=json.loads(Path(args.evidence).read_text())
    if proof.get("origin")!="native_desktop_screen_capture" or proof.get("is_pptx_playback_recording") is not True:
        raise ValueError("Use desktop playback evidence; design previews are not accepted")
    if proof["recording_sha256"]!=sha(video) or proof["source_pptx_sha256"]!=sha(deck):
        raise ValueError("Recording or PPTX differs from its evidence")
    metadata=json.loads(subprocess.check_output(["ffprobe","-v","error","-select_streams","v:0",
        "-show_entries","stream=width,height","-show_entries","format=duration","-of","json",str(video)],text=True))
    stream=metadata["streams"][0];vw,vh=stream["width"],stream["height"]
    if max(args.end,args.hold)>=float(metadata["format"]["duration"]):raise ValueError("Requested time is outside recording")
    prs=Presentation(deck)
    if abs(prs.slide_width/prs.slide_height-vw/vh)>.01:
        raise ValueError("Use a capture matching the slide aspect ratio; letterbox calibration is not implemented")
    slide_number=getattr(args,"slide",1)
    if not 1<=slide_number<=len(prs.slides):raise ValueError("Slide is outside the source PPTX")
    objects=getattr(args,"objects",None)
    bars=[]
    for shape in prs.slides[slide_number-1].shapes:
        selected=shape.name.lstrip("!") in objects if objects else args.prefix in shape.name
        if selected:
            bars.append({"name":shape.name,"x":round((shape.left+shape.width/2)*vw/prs.slide_width),
                         "top":round(shape.top*vh/prs.slide_height),
                         "bottom":round((shape.top+shape.height)*vh/prs.slide_height),
                         "color":tuple(bytes(shape.fill.fore_color.rgb))})
    if not bars:raise ValueError("No named solid-fill bars on the selected slide")
    if objects and len(bars)!=len(objects):raise ValueError("A requested native bar is missing")
    out.mkdir(parents=True,exist_ok=True);held=out/"held-reference.png"
    subprocess.run(["ffmpeg","-hide_banner","-loglevel","error","-y","-ss",str(args.hold),"-i",str(video),"-frames:v","1",str(held)],check=True)
    reference=Image.open(held).convert("RGB")
    pts=json.loads(subprocess.check_output(["ffprobe","-v","error","-select_streams","v:0",
        "-show_entries","frame=best_effort_timestamp_time","-of","json",str(video)],text=True))
    selected=[(i,float(f["best_effort_timestamp_time"])) for i,f in enumerate(pts["frames"])
              if args.start<=float(f["best_effort_timestamp_time"])<args.end]
    if not selected:raise ValueError("No captured frames in requested interval")
    times=[t for i,t in selected]
    y0=max(0,min(b["top"] for b in bars)-32);y1=min(vh,max(b["bottom"] for b in bars)+24)
    width,height=10,y1-y0;summary=[];traces=[];errors=[]
    for bar in bars:
        x0=bar["x"]-width//2;box=color_bounds(reference.crop((x0,y0,x0+width,y1)),bar["color"])
        if not box:raise ValueError("No held bar pixels: "+bar["name"])
        expected=box[3]+y0-1;records=[];positions=set();max_bottom=0;max_backward=0;last_top=None
        select=f"select=between(n\\,{selected[0][0]}\\,{selected[-1][0]}),crop={width}:{height}:{x0}:{y0}:exact=1"
        command=["ffmpeg","-hide_banner","-loglevel","error","-i",str(video),"-to",str(args.end+.25),"-vf",select,
                 "-fps_mode","passthrough","-f","rawvideo","-pix_fmt","rgb24","pipe:1"]
        process=subprocess.Popen(command,stdout=subprocess.PIPE);index=0
        try:
            while True:
                raw=process.stdout.read(width*height*3)
                if not raw:break
                if len(raw)!=width*height*3:raise RuntimeError("Partial decoded video frame")
                if index>=len(times):raise RuntimeError("Decoded frames do not match actual timestamps")
                bounds=color_bounds(Image.frombytes("RGB",(width,height),raw),bar["color"])
                if bounds and bounds[3]-bounds[1]>=3:
                    top,bottom=bounds[1]+y0,bounds[3]+y0-1
                    max_bottom=max(max_bottom,abs(bottom-expected))
                    if last_top is not None:max_backward=max(max_backward,top-last_top)
                    last_top=top;positions.add(top)
                    records.append({"timestamp":times[index],"top":top,"bottom":bottom})
                index+=1
            if process.wait()!=0 or index!=len(times):raise RuntimeError("Video decode or timestamp count failed")
        finally:
            if process.poll() is None:process.terminate();process.wait()
            process.stdout.close()
        summary.append({"name":bar["name"],"held_bottom_px":expected,"visible_frames":len(records),
                        "distinct_top_positions":len(positions),"max_bottom_deviation_px":max_bottom,
                        "max_backward_top_step_px":max_backward})
        traces.append({"bar":bar["name"],"observations":records})
        if max_bottom>args.tolerance or max_backward>args.tolerance:errors.append("Unstable column "+bar["name"])
        if len(positions)<5:errors.append("Insufficient reveal motion captured: "+bar["name"])
    report={"origin":"measurement_of_actual_desktop_recording","recording_sha256":sha(video),
            "source_pptx_sha256":sha(deck),"slide":slide_number,"time_interval_seconds":[args.start,args.end],
            "frames_sampled_per_bar":len(times),"held_reference_seconds":args.hold,"tolerance_px":args.tolerance,
            "summary":summary,"errors":errors,"records":traces,
            "scope":"Decoded native playback pixels, calibrated from held video frame. Measures color bounds in captured frames; does not infer missing frames or cross-player compatibility"}
    (out/"native-bar-playback-audit.json").write_text(json.dumps(report,indent=2)+"\n")
    return report


if __name__=="__main__":
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("video");parser.add_argument("--pptx",required=True);parser.add_argument("--evidence",required=True)
    parser.add_argument("--out",required=True);parser.add_argument("--prefix",default="comparison-bar-")
    parser.add_argument("--slide",type=int,default=1);parser.add_argument("--objects",nargs="+")
    parser.add_argument("--start",type=float,required=True);parser.add_argument("--end",type=float,required=True)
    parser.add_argument("--hold",type=float,required=True);parser.add_argument("--tolerance",type=float,default=1)
    result=check(parser.parse_args());print(json.dumps({k:v for k,v in result.items() if k!="records"},indent=2))
    if result["errors"]:raise SystemExit(1)
