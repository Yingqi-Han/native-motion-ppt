"""Evidence-based package checks and real static Office rendering."""
import hashlib,json,re,subprocess,tempfile,zipfile
from pathlib import Path
from lxml import etree
from pptx import Presentation
from PIL import Image,ImageDraw
from .pptx_export import NS,P,MC

def audit(path, strict=True):
    issues=[];slides=[]
    with zipfile.ZipFile(path) as z:
        if z.testzip():issues.append("ZIP checksum failure")
        for name in z.namelist():
            if name.endswith((".xml",".rels")):
                try:etree.fromstring(z.read(name))
                except etree.XMLSyntaxError:issues.append("Invalid XML: "+name)
        media=[n for n in z.namelist() if n.startswith("ppt/media/")]
        video=[n for n in media if n.lower().endswith((".mp4",".mov",".webm",".avi",".gif"))]
        if video:issues.append("Video/GIF media present")
        if strict and media:issues.append("Native-only deck contains media")
    p=Presentation(path)
    for i,s in enumerate(p.slides):
        root=s._element;ids=root.xpath(".//p:cNvPr/@id");tids=root.xpath(".//p:cTn/@id")
        if len(ids)!=len(set(ids)):issues.append(f"Slide {i+1}: duplicate object IDs")
        if len(tids)!=len(set(tids)):issues.append(f"Slide {i+1}: duplicate timing IDs")
        if any(v not in ids for v in root.xpath(".//p:spTgt/@spid")):issues.append(f"Slide {i+1}: broken animation target")
        names=root.xpath(".//p:cNvPr/@name")
        if len(names)!=len(set(names)):issues.append(f"Slide {i+1}: duplicate object names")
        if root.find(".//p:video",NS) is not None:issues.append(f"Slide {i+1}: video object")
        tags=[etree.QName(c).localname for c in root]
        if "timing" in tags and "extLst" in tags and tags.index("timing")>tags.index("extLst"):
            issues.append(f"Slide {i+1}: timing after extLst")
        count={k:len(root.findall(".//p:"+k,NS)) for k in ("sp","grpSp","pic","animEffect","animMotion","animRot","animScale","iterate")}
        clicks=len(root.xpath('.//p:cTn[@nodeType="clickEffect"]'))
        loops=len(root.xpath('.//p:cTn[@repeatCount="indefinite"]'))
        wipes=len(root.xpath('.//p:animEffect[starts-with(@filter,"wipe(")]'))
        slides.append({"slide":i+1,**count,"wipe_effects":wipes,"click_groups":clicks,"loop_effects":loops,"morph":len(root.findall(".//p159:morph",NS))})
    return {"file":str(Path(path).resolve()),"sha256":hashlib.sha256(Path(path).read_bytes()).hexdigest(),
            "native_only":strict,"media_parts":len(media),"slides":slides,"issues":issues,
            "scope":"Package, XML, object identities and animation references. Not a Microsoft PowerPoint playback test."}

def render(path,output,office_executable="soffice",office_env=None):
    import fitz
    output=Path(output);output.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="native-motion-office-") as tmp:
        profile=Path(tmp).as_uri()
        subprocess.run([office_executable,"-env:UserInstallation="+profile,"--headless","--convert-to","pdf:impress_pdf_Export","--outdir",str(output),str(Path(path).resolve())],
                       check=True,capture_output=True,text=True,timeout=180,env=office_env)
    pdf=output/(Path(path).stem+".pdf")
    if not pdf.exists():raise RuntimeError("Office renderer did not produce a PDF")
    doc=fitz.open(pdf);rows=[]
    thumb_h=round(800*doc[0].rect.height/doc[0].rect.width);row_h=thumb_h+30
    contact=Image.new("RGB",(1600,((len(doc)+1)//2)*row_h),(15,18,26));draw=ImageDraw.Draw(contact)
    prs=Presentation(path);missing=[];overflow=[];textbox_overflow=[]
    compact=lambda text:re.sub(r"\s+","",text)
    for i,page in enumerate(doc):
        page.get_pixmap(matrix=fitz.Matrix(1600/page.rect.width,1600/page.rect.width),alpha=False).save(output/f"slide-{i+1:02}.png")
        img=Image.open(output/f"slide-{i+1:02}.png");x=i%2*800;y=i//2*row_h
        contact.paste(img.resize((800,thumb_h)),(x,y));draw.text((x+18,y+thumb_h+6),f"{i+1:02} / NATIVE MOTION",fill=(168,177,197))
        actual=compact(page.get_text())
        source=prs.slides[i]._element.xpath(".//a:t/text()")
        for t in source:
            if compact(t) and compact(t) not in actual:missing.append({"slide":i+1,"text":t})
        for b in page.get_text("dict")["blocks"]:
            for line in b.get("lines",[]):
                for span in line["spans"]:
                    a,b,c,d=span["bbox"]
                    if a<-.5 or b<-.5 or c>page.rect.width+.5 or d>page.rect.height+.5:
                        overflow.append({"slide":i+1,"text":span["text"]})
        candidates=[]
        for sh in prs.slides[i].shapes:
            if getattr(sh,"has_text_frame",False) and sh.text.strip() and sh.rotation==0:
                factor=page.rect.width/prs.slide_width
                candidates.append((sh,compact(sh.text),(sh.left*factor,sh.top*factor,(sh.left+sh.width)*factor,(sh.top+sh.height)*factor)))
        from .fonts import glyph_ink
        for block in page.get_text("rawdict")["blocks"]:
            for line in block.get("lines",[]):
                text=compact("".join(ch["c"] for span in line["spans"] for ch in span["chars"]))
                if not text:continue
                box=line["bbox"]
                matches=[(sh,bounds) for sh,source,bounds in candidates if text in source]
                if matches:
                    def distance(pair):
                        a,b,c,d=pair[1]
                        dx=max(a-box[2],box[0]-c,0)
                        dy=max(b-box[3],box[1]-d,0)
                        centers=((box[0]+box[2]-a-c)/2)**2+((box[1]+box[3]-b-d)/2)**2
                        return dx*dx+dy*dy,centers
                    sh,bounds=min(matches,key=distance)
                    runs=[r for p in sh.text_frame.paragraphs for r in p.runs]
                    ink=[]
                    for span in line["spans"]:
                        for ch in span["chars"]:
                            metric=next((m for r in runs if r.font.name and (m:=glyph_ink(r.font.name,bool(r.font.bold),ch["c"])) and m[2] and m[2].replace("-","").casefold()==span["font"].replace("-","").casefold()),None)
                            if metric and metric[0]:
                                a,b,c,d=metric[0];scale=span["size"]/metric[1];ox,oy=ch["origin"]
                                ink.append((ox+a*scale,oy-d*scale,ox+c*scale,oy-b*scale))
                    if ink:
                        actual=(min(v[0] for v in ink),min(v[1] for v in ink),max(v[2] for v in ink),max(v[3] for v in ink))
                        if actual[0]<bounds[0]-1.5 or actual[1]<bounds[1]-1.5 or actual[2]>bounds[2]+1.5 or actual[3]>bounds[3]+1.5:
                            textbox_overflow.append({"slide":i+1,"object":sh.name,"text":text,"actual_ink":list(actual),"declared":list(bounds)})
        rows.append({"slide":i+1,"text":page.get_text(),"fonts":sorted({sp["font"] for b in page.get_text("dict")["blocks"] for l in b.get("lines",[]) for sp in l["spans"]})})
    contact.save(output/"contact.jpg",quality=93)
    return {"pages":len(doc),"missing_text":missing,"text_outside_canvas":overflow,"text_outside_declared_box":textbox_overflow,"slides":rows,
            "scope":"Actual LibreOffice pages; textbox ink checks cover nonrotated top-level text with a resolved matching font. Does not establish PowerPoint playback equivalence."}

def import_check(path,output,office_executable="soffice",office_env=None):
    output=Path(output);output.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="native-motion-import-") as tmp:
        subprocess.run([office_executable,"-env:UserInstallation="+Path(tmp).as_uri(),"--headless","--convert-to","odp","--outdir",str(output),str(Path(path).resolve())],
                       check=True,capture_output=True,text=True,timeout=180,env=office_env)
    with zipfile.ZipFile(output/(Path(path).stem+".odp")) as z:
        root=etree.fromstring(z.read("content.xml"))
    ns={"anim":"urn:oasis:names:tc:opendocument:xmlns:animation:1.0",
        "smil":"urn:oasis:names:tc:opendocument:xmlns:smil-compatible:1.0"}
    counts={k:len(root.findall(".//anim:"+k,ns)) for k in ("animateMotion","animateTransform","iterate","set","transitionFilter")}
    loops=sum(1 for n in root.iter() if n.get("{"+ns["smil"]+"}repeatCount")=="indefinite")
    filters=[{etree.QName(k).localname:v for k,v in n.attrib.items()} for n in root.findall('.//anim:transitionFilter',ns)]
    return {"counts":counts,"loop_effects":loops,"transition_filters":filters,"scope":"Actual Impress import counts; not a Microsoft PowerPoint slideshow test."}

