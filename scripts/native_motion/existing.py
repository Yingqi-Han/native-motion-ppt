"""Inspectable plans for existing PPTX; layout rules propose, never certify semantics."""
import hashlib
import html
import io
import json
import math
from pathlib import Path
import zipfile

from lxml import etree
from pptx import Presentation

P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
NS = {'p': P}


def digest(data):
    return hashlib.sha256(data).hexdigest()


def _shape(shape):
    sizes = []
    if shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            if paragraph.font.size:
                sizes.append(paragraph.font.size.pt)
            sizes += [run.font.size.pt for run in paragraph.runs if run.font.size]
    row = {'id': shape.shape_id, 'name': shape.name, 'kind': str(shape.shape_type),
           'native_type': shape.shape_type.name, 'text': getattr(shape, 'text', ''),
           'box_emu': [int(shape.left), int(shape.top), int(shape.width), int(shape.height)],
           'max_font_size_pt': max(sizes, default=None), 'rotation': float(shape.rotation),
           'placeholder': shape.placeholder_format.type.name if shape.is_placeholder else None,
           'is_chart': bool(getattr(shape, 'has_chart', False)), 'is_table': bool(shape.has_table)}
    if hasattr(shape, 'shapes'):
        row['children'] = [_shape(child) for child in shape.shapes]
    return row


def inspect_existing(source):
    source = Path(source)
    raw = source.read_bytes()
    prs = Presentation(io.BytesIO(raw))
    return {'format_version': 2, 'file': str(source.resolve()), 'source_pptx_sha256': digest(raw),
            'canvas_emu': [int(prs.slide_width), int(prs.slide_height)],
            'slides': [{'slide': i, 'part': str(slide.part.partname).lstrip('/'),
                        'has_animation': slide._element.find('p:timing', NS) is not None,
                        'objects': [_shape(shape) for shape in slide.shapes]}
                       for i, slide in enumerate(prs.slides, 1)]}


def _contains(outer, inner, pad=0):
    x, y, w, h = outer
    a, b, c, d = inner
    return a >= x-pad and b >= y-pad and a+c <= x+w+pad and b+d <= y+h+pad


def _role(obj, canvas, heading):
    w, h = canvas
    x, y, bw, bh = obj['box_emu']
    placeholder = obj['placeholder'] or ''
    if placeholder in {'FOOTER', 'DATE', 'SLIDE_NUMBER'}:
        return 'context', 'Native footer/date/page-number placeholder'
    if obj['id'] == heading:
        return 'title', 'Native title placeholder or short prominent text in the top area'
    if obj['text'].strip() and y < .14*h and obj['max_font_size_pt'] and obj['max_font_size_pt'] <= 18:
        return 'context', 'Small top text; proposed persistent header/page label, check actual meaning'
    if bw <= 0 or bh <= 0 or not _contains([0, 0, w, h], obj['box_emu'], w*.005):
        return 'manual', 'Zero-sized or off-canvas object needs visual review'
    if obj['rotation']:
        return 'manual', 'Rotated object: keep unchanged until its grouping is reviewed'
    if y >= .88*h and obj['text'] and (obj['max_font_size_pt'] or 0) <= 14:
        return 'context', 'Small bottom text; proposed persistent caption/footer, check actual meaning'
    if obj.get('children'):
        return 'visual', 'Existing native group remains one atomic target; child IDs are not flattened'
    if obj['is_chart'] or obj['is_table']:
        return 'data', 'Native chart/table reveals as a whole; axes/data geometry are not moved'
    if obj['text'].strip():
        return 'body', 'Whole text box; paragraph content and formatting stay unchanged'
    if obj['native_type'] in {'LINE', 'CONNECTOR'} or min(bw/w, bh/h) < .006 or bw*bh > .65*w*h:
        return 'context', 'Thin rule or large background; keep reading context visible'
    if obj['native_type'] in {'AUTO_SHAPE', 'FREEFORM', 'PICTURE'}:
        return 'visual', 'Whole native shape/image; source picture bytes remain unchanged'
    return 'manual', 'Unsupported object type; no guessed animation is applied'


def make_plan(source, max_clicks=3):
    if type(max_clicks) is not int or not 1 <= max_clicks <= 5:
        raise ValueError('max-clicks must be an integer from 1 to 5')
    inventory = inspect_existing(source)
    canvas = inventory['canvas_emu']
    plan = {'format_version': 2, 'source_pptx_sha256': inventory['source_pptx_sha256'],
            'planning_status': 'draft_requires_visual_review', 'max_clicks': max_clicks,
            'animation': {'effect': 'fade', 'duration_ms': 450, 'stagger_ms': 100},
            'slides': [], 'review': [],
            'scope': 'Layout-assisted draft, not a semantic or playback verdict. Inspect actual slides and edit groups before applying. Existing animated slides stay unchanged.'}
    width, height = canvas
    for slide in inventory['slides']:
        if slide['has_animation']:
            plan['review'].append({'slide': slide['slide'], 'status': 'preserved',
                                   'reason': 'Existing animation: entire slide is preserved byte-for-byte', 'objects': []})
            continue
        objects = slide['objects']
        titles = [o for o in objects if o['placeholder'] in {'TITLE', 'CENTER_TITLE', 'VERTICAL_TITLE'}]
        if not titles:
            candidates = [o for o in objects if o['text'].strip() and len(o['text']) <= 120
                          and o['box_emu'][1] < .28*height and o['max_font_size_pt']]
            if candidates:
                best = max(candidates, key=lambda o: o['max_font_size_pt'])
                other_sizes = [o['max_font_size_pt'] for o in objects if o != best and o['text'].strip() and o['max_font_size_pt']]
                if best['max_font_size_pt'] >= 24 and (not other_sizes or best['max_font_size_pt'] >= max(other_sizes)*1.2):
                    titles = [best]
        heading = titles[0]['id'] if titles else None
        reviewed = []
        for obj in objects:
            role, reason = _role(obj, canvas, heading)
            reviewed.append({**obj, 'role': role, 'reason': reason})
        eligible = [o for o in reviewed if o['role'] in {'body', 'visual', 'data'}]
        units = []
        used = set()
        # A card/container binds only fully contained content. Ambiguous overlaps
        # remain separate proposals; proximity alone never joins distant columns.
        for obj in sorted(eligible, key=lambda o: -o['box_emu'][2]*o['box_emu'][3]):
            if obj['id'] in used or obj['text'].strip() or obj.get('children') or obj['is_chart'] or obj['is_table']:
                continue
            inside = [other for other in eligible if other['id'] not in used and other['id'] != obj['id']
                      and _contains(obj['box_emu'], other['box_emu'])]
            if inside:
                unit = [obj, *sorted(inside, key=lambda o: (o['box_emu'][1], o['box_emu'][0]))]
                units.append(unit); used.update(o['id'] for o in unit)
        for obj in eligible:
            if obj['id'] not in used:
                units.append([obj]); used.add(obj['id'])
        units.sort(key=lambda unit: (round(min(o['box_emu'][1] for o in unit)/(height*.04)),
                                     min(o['box_emu'][0] for o in unit)))
        block = max(1, math.ceil(len(units)/max_clicks))
        groups = [[o['id'] for unit in units[i:i+block] for o in unit] for i in range(0, len(units), block)]
        if groups:
            plan['slides'].append({'slide': slide['slide'], 'groups': groups,
                                   'units': [[o['id'] for o in unit] for unit in units]})
        plan['review'].append({'slide': slide['slide'], 'status': 'proposed' if groups else 'unchanged',
                               'reason': 'Check argument order and keep chart annotations with their data; layout rules cannot infer the intended narrative',
                               'objects': reviewed})
    validate_plan(plan, inventory)
    return inventory, plan


def validate_plan(plan, inventory):
    if not isinstance(plan, dict) or not isinstance(plan.get('slides'), list):
        raise ValueError('Plan requires a slides list')
    version = plan.get('format_version', 1)
    if type(version) is not int or version not in (1, 2):
        raise ValueError('Unsupported plan format_version')
    if version == 2 and plan.get('source_pptx_sha256') != inventory['source_pptx_sha256']:
        raise ValueError('Plan source SHA-256 differs from this PPTX; regenerate the plan')
    if 'source_pptx_sha256' in plan and plan['source_pptx_sha256'] != inventory['source_pptx_sha256']:
        raise ValueError('Plan source SHA-256 differs from this PPTX')
    animation = plan.get('animation', {})
    if not isinstance(animation, dict) or animation.get('effect', 'fade') != 'fade':
        raise ValueError('Existing-PPT plans support a whole-object fade; use native authoring for other effects')
    for name, default, bounds in [('duration_ms', 450, (100, 3000)), ('stagger_ms', 100, (0, 1000))]:
        value = animation.get(name, default)
        if type(value) is not int or not bounds[0] <= value <= bounds[1]:
            raise ValueError('Invalid plan animation '+name)
    slides = {s['slide']: s for s in inventory['slides']}
    seen = set()
    for item in plan['slides']:
        if not isinstance(item, dict) or type(item.get('slide')) is not int or item['slide'] not in slides or item['slide'] in seen:
            raise ValueError('Plan has an invalid/duplicate slide number')
        seen.add(item['slide'])
        if not isinstance(item.get('groups'), list):
            raise ValueError('Plan groups must be a list')
        top_ids = {o['id'] for o in slides[item['slide']]['objects']};used = set()
        for group in item['groups']:
            if not isinstance(group, list) or not group:
                raise ValueError('Plan groups must be nonempty lists')
            for sid in group:
                if type(sid) is not int or sid not in top_ids or sid in used:
                    raise ValueError('Plan target is missing, repeated, or not a top-level object; use the parent group ID')
                used.add(sid)
    return {item['slide']: item['groups'] for item in plan['slides']}


def verify_changes(source, output, allowed_parts):
    changed = []
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(output) as after:
        if len(before.namelist()) != len(set(before.namelist())) or len(after.namelist()) != len(set(after.namelist())):
            raise ValueError('Duplicate ZIP members are not supported')
        if set(before.namelist()) != set(after.namelist()):
            raise ValueError('Enhancement added or removed package parts')
        for name in before.namelist():
            a, b = before.read(name), after.read(name)
            if a == b:
                continue
            if name not in allowed_parts:
                raise ValueError('Unrelated package part changed: '+name)
            old, new = etree.fromstring(a), etree.fromstring(b)
            for root in (old, new):
                for timing in root.findall('p:timing', NS):
                    root.remove(timing)
            if etree.tostring(old, method='c14n') != etree.tostring(new, method='c14n'):
                raise ValueError('Slide content outside animation timing changed: '+name)
            changed.append(name)
        return {'status': 'passed', 'source_pptx_sha256': digest(Path(source).read_bytes()),
                'output_pptx_sha256': digest(Path(output).read_bytes()), 'changed_timing_parts': sorted(changed),
                'unchanged_package_parts': len(before.namelist())-len(changed),
                'layout_and_relationships_preserved': True,
                'scope': 'ZIP member bytes and canonical slide XML outside p:timing; does not certify playback or visual quality'}


def write_plan_report(output, inventory, plan):
    output = Path(output);esc = html.escape
    width, height = inventory['canvas_emu'];pages = []
    for review in plan['review']:
        groups = next((s['groups'] for s in plan['slides'] if s['slide'] == review['slide']), [])
        order = {sid: i for i, group in enumerate(groups, 1) for sid in group}
        nodes = [];rows = []
        for obj in review['objects']:
            x,y,w,h = [v/width*1000 for v in obj['box_emu']]
            label = f"{order.get(obj['id'], '静态')} · {obj['id']} · {obj['name']}"
            nodes.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="none" stroke="'+('#00766a' if obj['id'] in order else '#9ba8a2')+'"/><text x="'+str(x+3)+'" y="'+str(y+16)+'">'+esc(label)+'</text>')
            rows.append(f'<tr><td>{obj["id"]}</td><td>{esc(obj["name"])}</td><td>{esc(obj["role"])}</td><td>{esc(obj["text"][:240])}</td><td>{esc(obj["reason"])}</td></tr>')
        pages.append(f'<section><h2>第 {review["slide"]} 页 · {esc(review["status"])}</h2><p>{esc(review["reason"])}</p><p>点击组：{esc(json.dumps(groups))}</p><svg viewBox="0 0 1000 {height/width*1000}">'+''.join(nodes)+'</svg><table><tr><th>ID</th><th>对象</th><th>建议角色</th><th>内容摘录</th><th>依据</th></tr>'+''.join(rows)+'</table></section>')
    (output/'plan-review.html').write_text('''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>现有 PPT 动画方案</title><style>body{max-width:1100px;margin:40px auto;padding:0 20px;background:#f4f2ea;color:#173330;font:15px/1.6 system-ui}svg{width:100%;background:white}svg text{font:14px system-ui;fill:#173330}table{border-collapse:collapse;width:100%}td,th{text-align:left;border-bottom:1px solid #ccd5cc;padding:10px;overflow-wrap:anywhere}section{margin:35px 0}code{overflow-wrap:anywhere}</style><h1>现有 PPT 动画方案</h1><p>这是布局辅助草案。框图显示原生对象的位置，不是幻灯片渲染或实播。先查看实际页面，核对标题、注释与讲解顺序，再编辑 plan.json 中的 groups；同组对象在同一次点击出现。</p><p>源文件 SHA-256：<code>'''+esc(inventory['source_pptx_sha256'])+'</code></p><p>标题和上下文默认静态；新动作仅为整体淡入，已有动画整页保留。未支持或模糊对象保持原样。</p>'+''.join(pages)+'</html>', encoding='utf-8')
