"""v1.1 regressions: editable curves, readable timing, fonts and safe upgrades."""
import copy,json,sys,tempfile,unittest
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/"scripts"))
from native_motion.model import normalize
from native_motion.geometry import parse_path
from native_motion.pptx_export import build
from native_motion.review import review,sample_times
from native_motion.choreography import choreograph
from native_motion.fonts import resolve_font
from install_skill import install
from pptx import Presentation
class QualityTests(unittest.TestCase):
 def doc(self):
  return {"version":1,"slides":[{"duration":4000,"objects":[
   {"id":"title","kind":"text","text":"把变化讲清楚。","box":[100,200,1200,180],"size":100,"role":"title"},
   {"id":"body","kind":"text","text":"先看条件，再看结果。","box":[100,650,1200,100],"size":40,"role":"body"}]}]}
 def test_native_curve_is_editable_without_media(self):
  d=self.doc();d["slides"][0]["objects"].append({"id":"curve","kind":"path","box":[1200,50,600,300],"fill":"none","stroke":"4148FF","d":"M 10 10 C 80 0 150 280 250 200 Q 400 100 500 240"})
  with tempfile.TemporaryDirectory() as t:
   p=Path(t)/"curve.pptx";build(normalize(d),p);s=Presentation(p).slides[0]
   self.assertEqual(len(s._element.xpath(".//a:cubicBezTo")),1)
   self.assertEqual(len(s._element.xpath(".//a:quadBezTo")),1)
   self.assertFalse(s._element.xpath(".//p:pic"))
 def test_unsupported_svg_commands_are_rejected(self):
  with self.assertRaises(ValueError):parse_path("M 0 0 A 10 10 0 0 0 30 30")
 def test_relative_and_repeated_move_coordinates(self):
  p=parse_path("m 10 20 30 40 c 1 2 3 4 5 6 z")
  self.assertEqual(p[1],{"op":"L","points":[[40.,60.]]})
  self.assertEqual(p[2]["points"][-1],[45.,66.])
 def test_fixed_leading_matches_design_pixels(self):
  with tempfile.TemporaryDirectory() as t:
   p=Path(t)/"deck.pptx";build(normalize(self.doc()),p)
   leading=Presentation(p).slides[0].shapes[0].text_frame.paragraphs[0].line_spacing
   self.assertAlmostEqual(leading.pt,59)
 def test_transition_is_included_in_cutoff_check(self):
  d=self.doc();d["slides"].append({"duration":2000,"transition":{"kind":"fade","duration":500},"objects":[{"id":"late","kind":"text","text":"稍后出现","box":[100,200,1000,120],"animations":[{"kind":"enter","at":1000,"duration":700}]}]})
  r=review(normalize(d));self.assertTrue(any(x["code"]=="animation-cut-off" for x in r["issues"]))
 def test_none_transition_does_not_delay_content(self):
  d=self.doc();d["slides"][0]["transition"]={"kind":"none","duration":999}
  self.assertEqual(normalize(d)["slides"][0]["transition"]["duration"],0)
 def test_late_actions_are_sampled(self):
  d=self.doc();d["slides"][0]["duration"]=6000
  d["slides"][0]["objects"][1]["animations"]=[{"kind":"enter","at":4000,"duration":500}]
  s=normalize(d)["slides"][0];times=sample_times(s,0)
  self.assertIn(4501,times);self.assertIn(5999,times)
 def test_roles_compile_to_clicks_and_preserve_authored_actions(self):
  d=self.doc();d["slides"][0]["objects"][0]["animations"]=[{"kind":"enter","at":200,"duration":600}]
  n,r=choreograph(d,"click",True)
  self.assertEqual(n["slides"][0]["objects"][0]["animations"][0]["at"],200)
  self.assertEqual(n["slides"][0]["objects"][1]["animations"][0]["click"],1)
  self.assertGreaterEqual(review(n)["slides"][0]["hold_ms"],900)
 def test_low_contrast_is_diagnosed(self):
  d=self.doc();d["slides"][0]["objects"][0]["fill"]="12141A"
  self.assertTrue(any(x["code"]=="text-contrast" for x in review(normalize(d))["issues"]))
 def test_missing_font_is_not_silently_verified(self):
  info=resolve_font("Native Motion No Such Font")
  self.assertTrue(info is None or not info["matched"])
 def test_nonfinite_style_value_is_rejected_before_render(self):
  d=self.doc();d["slides"][0]["objects"][0]["spacing"]=float("nan")
  with self.assertRaises(ValueError):normalize(d)
 def test_chained_cubic_motion_survives_native_export(self):
  d=self.doc();d["slides"][0]["objects"][0]["animations"]=[{"kind":"move","path":[[0,0],[10,0],[20,10],[20,20],[20,30],[10,40],[0,40]],"duration":1500}]
  with tempfile.TemporaryDirectory() as t:
   p=Path(t)/"curve.pptx";build(normalize(d),p)
   path=Presentation(p).slides[0]._element.xpath(".//p:animMotion")[0].get("path")
   self.assertEqual(path.count("C "),2)
   self.assertTrue(path.endswith(" E"))
 def test_incomplete_cubic_chain_is_rejected(self):
  d=self.doc();d["slides"][0]["objects"][0]["animations"]=[{"kind":"move","path":[[0,0]]*5}]
  with self.assertRaises(ValueError):normalize(d)
 def test_loop_defaults_use_symmetric_acceleration(self):
  d=self.doc();d["slides"][0]["objects"][0]["animations"]=[{"kind":"rotate","loop":True}]
  e=normalize(d)["slides"][0]["objects"][0]["animations"][0]
  self.assertEqual((e["accel"],e["decel"]),(.5,.5))
 def test_ambient_motion_does_not_reduce_reading_hold_but_cutoff_is_reported(self):
  d=self.doc();d["slides"][0]["objects"].append({"id":"orb","role":"visual","kind":"ellipse","box":[1500,200,200,200],"animations":[{"kind":"move","duration":3800,"to":[0,-12],"ambient":True}]})
  r=review(normalize(d));self.assertEqual(r["slides"][0]["hold_ms"],4000)
  d["slides"][0]["objects"][-1]["animations"][0]["duration"]=4500
  self.assertTrue(any(i["code"]=="animation-cut-off" for i in review(normalize(d))["issues"]))
 def test_content_cannot_bypass_reading_hold_as_ambient_entry(self):
  d=self.doc();d["slides"][0]["objects"][0]["animations"]=[{"kind":"enter","ambient":True}]
  with self.assertRaises(ValueError):normalize(d)
 def test_managed_upgrade_preserves_extra_files_and_backup(self):
  with tempfile.TemporaryDirectory() as t:
   t=Path(t);src=t/"src";src.mkdir();(src/"VERSION").write_text("1.0.0");(src/"SKILL.md").write_text("old")
   result=install(src,t/"workspace");dest=Path(result["path"]);(dest/"user.txt").write_text("keep")
   (src/"VERSION").write_text("1.1.0");(src/"SKILL.md").write_text("new")
   upgraded=install(src,t/"workspace",True)
   self.assertEqual((dest/"user.txt").read_text(),"keep")
   self.assertEqual((Path(upgraded["backup"])/"SKILL.md").read_text(),"old")
 def test_managed_upgrade_refuses_modified_tracked_file(self):
  with tempfile.TemporaryDirectory() as t:
   t=Path(t);src=t/"src";src.mkdir();(src/"SKILL.md").write_text("original")
   dest=Path(install(src,t/"workspace")["path"]);(dest/"SKILL.md").write_text("user edit")
   (src/"SKILL.md").write_text("new version")
   with self.assertRaises(ValueError):install(src,t/"workspace",True)
   self.assertEqual((dest/"SKILL.md").read_text(),"user edit")
if __name__=="__main__":unittest.main()

