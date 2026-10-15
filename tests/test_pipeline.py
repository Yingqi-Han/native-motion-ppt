"""Behavioral regression checks for native output and existing-deck preservation."""
import copy,json,tempfile,unittest,sys,zipfile
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/"scripts"))
from native_motion.model import normalize
from native_motion.pptx_export import build,enhance
from native_motion.qa import audit
from pptx import Presentation
from pptx.util import Inches,Pt
from lxml import etree
class PipelineTests(unittest.TestCase):
 def doc(self):
  return {"version":1,"slides":[{"duration":3000,"clicks":[800],"objects":[
    {"id":"title","kind":"text","text":"真实对象","size":60,"box":[100,100,600,100]},
    {"id":"body","kind":"text","text":"分步展示","size":42,"box":[100,350,600,100],
     "animations":[{"kind":"enter","click":1,"duration":450,"by_letter":True}]}]}]}
 def test_missing_click_is_rejected(self):
  d=self.doc();d["slides"][0]["clicks"]=[]
  with self.assertRaises(ValueError):normalize(d)
 def test_duplicate_identity_is_rejected(self):
  d=self.doc();d["slides"][0]["objects"][1]["id"]="title"
  with self.assertRaises(ValueError):normalize(d)
 def test_native_target_and_click_group(self):
  with tempfile.TemporaryDirectory() as t:
   p=Path(t)/"deck.pptx";build(normalize(self.doc()),p)
   r=audit(p);self.assertEqual(r["issues"],[]);self.assertEqual(r["media_parts"],0)
   self.assertEqual(r["slides"][0]["click_groups"],1);self.assertEqual(r["slides"][0]["iterate"],1)
 def test_auto_deck_converts_scheduled_clicks(self):
  with tempfile.TemporaryDirectory() as t:
   p=Path(t)/"auto.pptx";build(normalize(self.doc()),p,auto=True)
   r=audit(p);self.assertEqual(r["issues"],[]);self.assertEqual(r["slides"][0]["click_groups"],0)
 def test_semantic_plan_targets_the_requested_objects(self):
  with tempfile.TemporaryDirectory() as t:
   p=Path(t)/"source.pptx";q=Path(t)/"result.pptx"
   d=self.doc();d["slides"][0]["objects"][1]["animations"]=[]
   build(normalize(d),p);src=Presentation(p);ids=[sh.shape_id for sh in src.slides[0].shapes]
   result=enhance(p,q,plan={"slides":[{"slide":1,"groups":[ids]}]})
   self.assertEqual(result[0]["added"],2);self.assertEqual(result[0]["click_groups"],1)
   self.assertEqual(etree.tostring(src.slides[0]._element.spTree),etree.tostring(Presentation(q).slides[0]._element.spTree))
 def test_enhance_keeps_existing_timing_and_layout(self):
  with tempfile.TemporaryDirectory() as t:
   p=Path(t)/"source.pptx";q=Path(t)/"result.pptx";build(normalize(self.doc()),p)
   result=enhance(p,q)
   self.assertIn("skipped",result[0])
   with zipfile.ZipFile(p) as z,zipfile.ZipFile(q) as w:
    self.assertEqual(z.read("ppt/slides/slide1.xml"),w.read("ppt/slides/slide1.xml"))
 def test_enhance_static_deck_preserves_geometry_and_text(self):
  with tempfile.TemporaryDirectory() as t:
   p=Path(t)/"source.pptx";q=Path(t)/"result.pptx";prs=Presentation();s=prs.slides.add_slide(prs.slide_layouts[6])
   for i in range(5):
    sh=s.shapes.add_textbox(Inches(.5),Inches(.5+i),Inches(6),Inches(.7))
    sh.text="Heading" if i==0 else f"Content {i}";sh.text_frame.paragraphs[0].runs[0].font.size=Pt(36 if i==0 else 20)
   prs.save(p);enhance(p,q,clicks=2)
   before=Presentation(p);after=Presentation(q)
   self.assertEqual(etree.tostring(before.slides[0]._element.spTree),etree.tostring(after.slides[0]._element.spTree))
   r=audit(q);self.assertEqual(r["issues"],[]);self.assertEqual(r["slides"][0]["click_groups"],2)
if __name__=="__main__":unittest.main()

