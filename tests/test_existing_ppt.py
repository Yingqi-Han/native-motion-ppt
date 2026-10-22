"""Existing-deck safety and useful grouping behavior. Fixtures are not WPS evidence."""
import copy
import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1];sys.path.insert(0, str(ROOT/'scripts'))
from native_motion.existing import inspect_existing, make_plan, validate_plan, verify_changes, write_plan_report
from native_motion.model import normalize
from native_motion.pptx_export import build, enhance


class ExistingPptTests(unittest.TestCase):
    def fixture(self, root):
        path = root/'source.pptx';prs = Presentation();slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = 'Title stays visible'
        slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(36)
        for i in range(2):
            card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.5+i*4.7), Inches(2.0), Inches(4), Inches(2.6));card.name = 'Card '+str(i)
            for j in range(2):
                text = slide.shapes.add_textbox(Inches(.7+i*4.7), Inches(2.2+j*.8), Inches(3.6), Inches(.6))
                text.text = ['Argument', 'Supporting evidence'][j]+str(i)
                text.text_frame.paragraphs[0].runs[0].font.size = Pt(24 if j == 0 else 18)
        footer = slide.shapes.add_textbox(Inches(.5), Inches(7), Inches(8), Inches(.3));footer.text = 'Simulated source, not a customer project'
        footer.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
        slide.notes_slide.notes_text_frame.text = 'Preserve the original speaker notes.'
        image = io.BytesIO();Image.new('RGB', (8, 8), '#00766a').save(image, format='PNG');image.seek(0)
        slide.shapes.add_picture(image, Inches(.5), Inches(5.5), Inches(.5), Inches(.5))
        other = prs.slides.add_slide(prs.slide_layouts[6]);group = other.shapes.add_group_shape()
        group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2), Inches(2), Inches(1), Inches(1))
        group.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3), Inches(2), Inches(1), Inches(1))
        prs.save(path)
        return path

    def test_cards_are_atomic_units_and_titles_footers_stay_visible(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = self.fixture(Path(tmp));inventory, plan = make_plan(source, 3)
            row = plan['slides'][0];objects = inventory['slides'][0]['objects']
            targets = {sid for group in row['groups'] for sid in group}
            self.assertNotIn(objects[0]['id'], targets)
            self.assertNotIn(next(o['id'] for o in objects if 'Simulated' in o['text']), targets)
            for i in range(2):
                card_ids = {o['id'] for o in objects if o['name'] == 'Card '+str(i) or o['text'].endswith(str(i))}
                self.assertTrue(any(card_ids <= set(group) for group in row['groups']))
            self.assertEqual(plan['source_pptx_sha256'], inventory['source_pptx_sha256'])
            self.assertEqual(plan['planning_status'], 'draft_requires_visual_review')

    def test_native_group_is_not_flattened(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = self.fixture(Path(tmp));inventory, plan = make_plan(source)
            group = inventory['slides'][1]['objects'][0]
            self.assertEqual(plan['slides'][1]['groups'], [[group['id']]])
            bad = copy.deepcopy(plan);bad['slides'][1]['groups'] = [[group['children'][0]['id']]]
            with self.assertRaisesRegex(ValueError, 'top-level'):
                validate_plan(bad, inventory)

    def test_small_page_header_is_not_a_click_target(self):
        with tempfile.TemporaryDirectory() as tmp:
            source=self.fixture(Path(tmp));prs=Presentation(source)
            label=prs.slides[0].shapes.add_textbox(Inches(.5),Inches(.1),Inches(5),Inches(.3));label.text='SECTION / 01'
            label.text_frame.paragraphs[0].runs[0].font.size=Pt(11);sid=label.shape_id;prs.save(source)
            _,plan=make_plan(source)
            self.assertNotIn(sid,{v for g in plan['slides'][0]['groups'] for v in g})

    def test_enhancement_only_changes_timing_and_preserves_unknown_parts(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp);source = self.fixture(root)
            with zipfile.ZipFile(source, 'a') as z:
                z.writestr('customXml/preserved-regression.bin', b'unknown fixture bytes\x00\xff')
            before = source.read_bytes();inventory, plan = make_plan(source)
            output = root/'enhanced.pptx';enhance(source, output, plan=plan)
            report = verify_changes(source, output, {s['part'] for s in inventory['slides']})
            self.assertEqual(source.read_bytes(), before)
            self.assertTrue(report['layout_and_relationships_preserved'])
            self.assertEqual(len(report['changed_timing_parts']), 2)
            with zipfile.ZipFile(source) as a, zipfile.ZipFile(output) as b:
                self.assertEqual(set(a.namelist()), set(b.namelist()))
                for name in a.namelist():
                    if name not in report['changed_timing_parts']:
                        self.assertEqual(a.read(name), b.read(name), name)
                for name in report['changed_timing_parts']:
                    tree = etree.fromstring(b.read(name))
                    self.assertFalse(tree.findall('.//{http://schemas.openxmlformats.org/presentationml/2006/main}animMotion'))

    def test_stale_plan_and_invalid_later_slide_write_no_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp);source = self.fixture(root);_,plan = make_plan(source)
            bad = copy.deepcopy(plan);bad['source_pptx_sha256'] = '0'*64
            for candidate in [bad, {**plan, 'slides': [*plan['slides'], {'slide': 100, 'groups': [[4]]}]}]:
                with self.assertRaises(ValueError):enhance(source, root/'result.pptx', plan=candidate)
                self.assertFalse((root/'result.pptx').exists())

    def test_boolean_duplicate_and_empty_targets_are_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = self.fixture(Path(tmp));inventory, plan = make_plan(source);sid=plan['slides'][0]['groups'][0][0]
            for groups in [[[True]], [[sid], [sid]], [[]]]:
                bad = copy.deepcopy(plan);bad['slides'][0]['groups'] = groups
                with self.assertRaises(ValueError):validate_plan(bad, inventory)

    def test_already_animated_slide_is_an_exact_zip_copy(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);source=root/'source.pptx'
            build(normalize({'version':1,'slides':[{'duration':3000,'objects':[{'id':'a','kind':'rect','box':[100,100,100,100],
                  'animations':[{'kind':'move','to':[20,0],'duration':1000}]}]}]}), source)
            _,plan=make_plan(source);self.assertEqual(plan['slides'],[])
            enhance(source,root/'after.pptx',plan=plan)
            with zipfile.ZipFile(source) as a,zipfile.ZipFile(root/'after.pptx') as b:
                for name in a.namelist():self.assertEqual(a.read(name),b.read(name),name)

    def test_existing_output_is_preserved_and_plan_html_escapes_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);source=self.fixture(root);inventory,plan=make_plan(source)
            output=root/'existing.pptx';output.write_bytes(b'user original')
            with self.assertRaisesRegex(ValueError,'existing files'):enhance(source,output,plan=plan)
            self.assertEqual(output.read_bytes(),b'user original')
            plan['review'][0]['objects'][0]['name']='<script>alert(1)</script>'
            write_plan_report(root,inventory,plan);text=(root/'plan-review.html').read_text()
            self.assertNotIn('<script>alert(1)</script>',text);self.assertIn('&lt;script&gt;',text)


if __name__=='__main__':unittest.main()
